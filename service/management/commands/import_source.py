from datetime import datetime
from pathlib import Path
from django.conf import settings
from django.contrib.auth.models import Group, User
from django.core.management import call_command
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from service.models import Directory, Machine, Maintenance, Claim


def text(value):
    return str(value).strip() if value is not None else ''


def item(kind, name):
    return Directory.objects.get_or_create(kind=kind, name=text(name), defaults={
        'description': 'Импортировано из исходной таблицы. Подробное описание уточняется у производителя.'})[0]


def account(name, group):
    user = User.objects.filter(first_name=text(name), groups__name=group).first()
    if user is None:
        prefix = 'client' if group == 'Клиент' else 'service'
        number = 1
        while User.objects.filter(username=f'{prefix}{number}').exists():
            number += 1
        user = User.objects.create_user(username=f'{prefix}{number}', first_name=text(name), password=None)
        user.groups.add(Group.objects.get(name=group))
    return user


class Command(BaseCommand):
    help = 'Импорт исходного Excel. Существующие машины и записи не перезаписываются.'

    def add_arguments(self, parser):
        parser.add_argument('path', nargs='?', default=str(settings.BASE_DIR / 'materials/source.xlsx'))

    @transaction.atomic
    def handle(self, *args, **options):
        if not Path(options['path']).exists():
            raise CommandError('Файл Excel не найден.')
        call_command('setup_roles')
        workbook = load_workbook(options['path'], data_only=True)
        for row in workbook.worksheets[0].iter_rows(min_row=4, values_only=True):
            if not row[2] or not isinstance(row[11], datetime):
                continue
            Machine.objects.get_or_create(serial_number=text(row[2]), defaults={
                'machine_model': item('machine', row[1]), 'engine_model': item('engine', row[3]),
                'engine_number': text(row[4]), 'transmission_model': item('transmission', row[5]),
                'transmission_number': text(row[6]), 'drive_axle_model': item('drive_axle', row[7]),
                'drive_axle_number': text(row[8]), 'steer_axle_model': item('steer_axle', row[9]),
                'steer_axle_number': text(row[10]), 'shipment_date': row[11].date(),
                'supply_contract': 'Не указан в исходных данных', 'consignee': text(row[13]),
                'delivery_address': text(row[14]), 'equipment': text(row[15]),
                'client': account(row[12], 'Клиент'), 'service_company': account(row[16], 'Сервисная организация'),
            })
        for row in workbook.worksheets[1].iter_rows(min_row=2, values_only=True):
            if not row[0] or not isinstance(row[2], datetime):
                continue
            machine = Machine.objects.get(serial_number=text(row[0]))
            Maintenance.objects.get_or_create(machine=machine, work_order_number=text(row[4]),
                maintenance_date=row[2].date(), defaults={
                    'maintenance_type': item('maintenance', row[1]), 'operating_hours': int(row[3]),
                    'work_order_date': row[5].date(), 'organization': item('organization', row[6]),
                    'service_company': machine.service_company,
                })
        for row in workbook.worksheets[2].iter_rows(min_row=3, values_only=True):
            if not row[0] or not isinstance(row[1], datetime):
                continue
            machine = Machine.objects.get(serial_number=text(row[0]))
            Claim.objects.get_or_create(machine=machine, failure_date=row[1].date(), description=text(row[4]), defaults={
                'operating_hours': int(row[2]), 'failure_node': item('failure', row[3]),
                'recovery_method': item('recovery', row[5]), 'spare_parts': text(row[6]),
                'recovery_date': row[7].date() if row[7] else None, 'service_company': machine.service_company,
            })
        self.stdout.write(self.style.SUCCESS(
            f'В базе: {Machine.objects.count()} машин, {Maintenance.objects.count()} ТО, {Claim.objects.count()} рекламаций.'))
